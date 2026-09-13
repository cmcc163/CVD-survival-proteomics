import os
from pathlib import Path
import math

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lifelines import KaplanMeierFitter
from lifelines.utils import concordance_index


PRED_DIR = Path(os.environ.get("CVD_PREDICTION_ROOT", "artifacts/predictions"))
SUMMARY_XLSX = Path(os.environ.get("CVD_SUMMARY_XLSX", "results/summary.xlsx"))
OUTPUT_DIR = Path(__file__).resolve().parent

OUTCOMES = ["Total_CVD", "ASCVD", "HF"]
OUTCOME_LABELS = {"Total_CVD": "Total CVD", "ASCVD": "ASCVD", "HF": "HF"}
MODEL_ORDER = [
    "PREVENT",
    "LinearModel",
    "XGBoost",
    "MLP",
    "TabNet",
    "NODE",
    "FT-Transformer",
    "SAINT",
    "TabPFN",
]
MODEL_LABELS = {
    "PREVENT": "PREVENT equation",
    "LinearModel": "Refitted Cox model",
    "XGBoost": "XGBoost",
    "MLP": "MLP",
    "TabNet": "TabNet",
    "NODE": "NODE",
    "FT-Transformer": "FT-Transformer",
    "SAINT": "SAINT",
    "TabPFN": "TabPFN",
}
PREDICTOR_SETS = [
    ("Clinical predictors only", "classic", "classic_val.csv", ["classic", "all_classic"]),
    ("Protein markers only", "protein_lassonet", "protein_lassonet_val.csv", ["protein_lassonet"]),
    (
        "Clinical predictors plus protein markers",
        "all_lassonet",
        "all_lassonet_val.csv",
        ["all_lassonet"],
    ),
]


def format_int(n: int) -> str:
    return f"{int(n):,}".replace(",", " ")


def format_metric(point: float, ci_text: str, multiline: bool = False) -> str:
    low, high = parse_ci(ci_text)
    if multiline:
        return f"{point:.4f}\n({low:.4f}-{high:.4f})"
    return f"{point:.4f} ({low:.4f}-{high:.4f})"


def parse_ci(ci_text: str) -> tuple[float, float]:
    if pd.isna(ci_text):
        return math.nan, math.nan
    parts = str(ci_text).replace("(", "").replace(")", "").split(",")
    if len(parts) != 2:
        return math.nan, math.nan
    return float(parts[0].strip()), float(parts[1].strip())


def protein_marker_counts() -> dict[str, int]:
    counts = {}
    for outcome in OUTCOMES:
        path = BASE_DIR / outcome / "lassonet_protein.csv"
        counts[outcome] = len(pd.read_csv(path))
    return counts


def no_protein_markers(config: str, counts: dict[str, int]) -> str:
    if config == "classic":
        return "0"
    return " / ".join(str(counts[outcome]) for outcome in OUTCOMES)


def load_summary() -> dict[str, pd.DataFrame]:
    return {sheet: pd.read_excel(SUMMARY_XLSX, sheet_name=sheet) for sheet in OUTCOMES}


def find_ci(
    summary: dict[str, pd.DataFrame],
    outcome: str,
    model: str,
    configs: list[str],
    metric: str,
    validation: str = "european",
) -> str:
    df = summary[outcome]
    matched = df[df["model"].eq(model) & df["config"].isin(configs)]
    if matched.empty:
        raise FileNotFoundError(f"No bootstrap CI found for {outcome}, {model}, {configs}, {metric}")
    return str(matched.iloc[0][f"{validation}_{metric}_CI"])


def prediction_file(model: str, outcome: str, filename: str) -> Path:
    path = PRED_DIR / model / outcome / filename
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def filename_for(config_key: str, validation: str) -> str:
    suffix = {"european": "val", "asian": "asian", "other": "other"}[validation]
    if config_key == "classic":
        return f"classic_{suffix}.csv"
    return f"{config_key}_{suffix}.csv"


def c_index_from_predictions(path: Path) -> float:
    data = pd.read_csv(path, usecols=["time", "event", "risk_score"])
    data = data.dropna(subset=["time", "event", "risk_score"])
    return float(concordance_index(data["time"], -data["risk_score"], data["event"]))


def censoring_survival_at(kmf: KaplanMeierFitter, times: np.ndarray) -> np.ndarray:
    survival = kmf.survival_function_[kmf.survival_function_.columns[0]]
    return np.asarray(survival.reindex(times, method="ffill").fillna(1.0), dtype=float)


def cumulative_dynamic_auc_10yr(path: Path) -> float:
    data = pd.read_csv(path, usecols=["time", "event", "prob_10yr"])
    data = data.dropna(subset=["time", "event", "prob_10yr"]).copy()
    data["time"] = pd.to_numeric(data["time"], errors="coerce")
    data["event"] = pd.to_numeric(data["event"], errors="coerce")
    data["prob_10yr"] = pd.to_numeric(data["prob_10yr"], errors="coerce")
    data = data.dropna()
    horizon = 365.25 * 10

    cases = data[(data["event"].eq(1)) & (data["time"] <= horizon)].copy()
    controls = data[data["time"] > horizon].copy()
    if cases.empty or controls.empty:
        return math.nan

    kmf = KaplanMeierFitter()
    kmf.fit(data["time"], event_observed=1 - data["event"])
    case_weights = 1.0 / np.clip(censoring_survival_at(kmf, cases["time"].to_numpy()), 1e-8, None)
    case_scores = cases["prob_10yr"].to_numpy()
    control_scores = controls["prob_10yr"].to_numpy()

    total = 0.0
    for score, weight in zip(case_scores, case_weights):
        total += weight * (
            np.sum(score > control_scores) + 0.5 * np.sum(np.isclose(score, control_scores))
        )
    return float(total / (np.sum(case_weights) * len(control_scores)))


def models_for_config(config: str) -> list[str]:
    if config == "classic":
        return MODEL_ORDER
    return [model for model in MODEL_ORDER if model != "PREVENT"]


def build_metric_rows(metric: str) -> pd.DataFrame:
    summary = load_summary()
    counts = protein_marker_counts()
    rows = []
    for group_label, config_key, filename, summary_configs in PREDICTOR_SETS:
        rows.append(
            {
                "Model": group_label,
                "No. of protein markers": "",
                "Total CVD": "",
                "ASCVD": "",
                "HF": "",
                "row_type": "section",
            }
        )
        for model in models_for_config(config_key):
            row = {
                "Model": MODEL_LABELS[model],
                "No. of protein markers": no_protein_markers(config_key, counts),
                "row_type": "data",
            }
            for outcome in OUTCOMES:
                path = prediction_file(model, outcome, filename)
                if metric == "C-index":
                    point = c_index_from_predictions(path)
                    ci = find_ci(summary, outcome, model, summary_configs, "C-index")
                elif metric == "AUC10yr":
                    point = cumulative_dynamic_auc_10yr(path)
                    ci = find_ci(summary, outcome, model, summary_configs, "AUC10yr")
                else:
                    raise ValueError(metric)
                row[OUTCOME_LABELS[outcome]] = format_metric(point, ci)
            rows.append(row)
    return pd.DataFrame(rows)


def build_metric_wide(metric: str, validation: str = "european") -> pd.DataFrame:
    summary = load_summary()
    counts = protein_marker_counts()
    rows = []
    for group_label, config_key, _, summary_configs in PREDICTOR_SETS:
        section = {
            "Outcome": group_label,
            "No. of protein markers": "",
            "row_type": "section",
        }
        for model in MODEL_ORDER:
            section[MODEL_LABELS[model]] = ""
        rows.append(section)

        for outcome in OUTCOMES:
            row = {
                "Outcome": OUTCOME_LABELS[outcome],
                "No. of protein markers": "/" if config_key == "classic" else str(counts[outcome]),
                "row_type": "data",
            }
            for model in MODEL_ORDER:
                if model not in models_for_config(config_key):
                    row[MODEL_LABELS[model]] = "/"
                    continue
                path = prediction_file(model, outcome, filename_for(config_key, validation))
                if metric == "C-index":
                    point = c_index_from_predictions(path)
                    ci = find_ci(summary, outcome, model, summary_configs, "C-index", validation)
                elif metric == "AUC10yr":
                    point = cumulative_dynamic_auc_10yr(path)
                    ci = find_ci(summary, outcome, model, summary_configs, "AUC10yr", validation)
                else:
                    raise ValueError(metric)
                row[MODEL_LABELS[model]] = format_metric(point, ci, multiline=True)
            rows.append(row)
    ordered_columns = [
        "Outcome",
        "No. of protein markers",
        *[MODEL_LABELS[model] for model in MODEL_ORDER],
        "row_type",
    ]
    return pd.DataFrame(rows)[ordered_columns]


def build_validation_wide(metric: str, validations: list[str]) -> pd.DataFrame:
    summary = load_summary()
    counts = protein_marker_counts()
    validation_labels = {
        "european": "European hold-out set",
        "asian": "Asian ancestry",
        "other": "Other ancestry",
    }
    rows = []
    for group_label, config_key, _, summary_configs in PREDICTOR_SETS:
        section = {
            "Validation set": group_label,
            "Outcome": "",
            "No. of protein markers": "",
            "row_type": "section",
        }
        for model in MODEL_ORDER:
            section[MODEL_LABELS[model]] = ""
        rows.append(section)

        for validation in validations:
            for outcome in OUTCOMES:
                row = {
                    "Validation set": validation_labels[validation],
                    "Outcome": OUTCOME_LABELS[outcome],
                    "No. of protein markers": "/" if config_key == "classic" else str(counts[outcome]),
                    "row_type": "data",
                }
                for model in MODEL_ORDER:
                    if model not in models_for_config(config_key):
                        row[MODEL_LABELS[model]] = "/"
                        continue
                    path = prediction_file(model, outcome, filename_for(config_key, validation))
                    if metric == "C-index":
                        point = c_index_from_predictions(path)
                        ci = find_ci(summary, outcome, model, summary_configs, "C-index", validation)
                    elif metric == "AUC10yr":
                        point = cumulative_dynamic_auc_10yr(path)
                        ci = find_ci(summary, outcome, model, summary_configs, "AUC10yr", validation)
                    else:
                        raise ValueError(metric)
                    row[MODEL_LABELS[model]] = format_metric(point, ci, multiline=True)
                rows.append(row)

    ordered_columns = [
        "Validation set",
        "Outcome",
        "No. of protein markers",
        *[MODEL_LABELS[model] for model in MODEL_ORDER],
        "row_type",
    ]
    return pd.DataFrame(rows)[ordered_columns]


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge, attrs in kwargs.items():
        element = tc_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def clear_borders(cell):
    nil = {"val": "nil"}
    set_cell_border(cell, top=nil, bottom=nil, left=nil, right=nil)


def style_paragraph(paragraph, font_size=8.5, bold=False, color=None):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(font_size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)


def set_landscape(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)


def set_cell_text(cell, text: str, font_size: float, bold: bool = False) -> None:
    cell.text = ""
    lines = str(text).split("\n")
    paragraph = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(font_size)
        run.bold = bold
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def write_wide_metric_docx(
    data: pd.DataFrame,
    path: Path,
    title_prefix: str,
    title: str,
    note: str,
) -> None:
    document = Document()
    set_landscape(document)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(title_prefix)
    run.bold = True
    run.font.color.rgb = RGBColor(192, 0, 0) if title_prefix.startswith("Table") else RGBColor(0, 0, 0)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)

    display = data.drop(columns=["row_type"])
    table = document.add_table(rows=1, cols=len(display.columns))
    first_width = Inches(1.05) if "Validation set" not in display.columns else Inches(1.2)
    second_width = Inches(1.0)
    marker_width = Inches(0.75)
    model_width = Inches(0.86)
    widths = []
    for column in display.columns:
        if column in {"Outcome"}:
            widths.append(first_width)
        elif column == "Validation set":
            widths.append(first_width)
        elif column == "No. of protein markers":
            widths.append(marker_width)
        elif column == "Outcome" and "Validation set" in display.columns:
            widths.append(second_width)
        else:
            widths.append(model_width)
    if "Validation set" in display.columns:
        widths[1] = Inches(0.85)

    for i, column in enumerate(display.columns):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "BFBFBF")
        clear_borders(cell)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "12", "color": "000000"},
            bottom={"val": "single", "sz": "8", "color": "000000"},
        )
        set_cell_text(cell, column, 6.7, bold=True)

    for idx, row in data.iterrows():
        cells = table.add_row().cells
        values = display.loc[idx].tolist()
        is_section = row["row_type"] == "section"
        for i, value in enumerate(values):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            clear_borders(cells[i])
            if is_section:
                set_cell_shading(cells[i], "F2F2F2")
            set_cell_text(cells[i], value, 6.5, bold=is_section)

    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom={"val": "single", "sz": "12", "color": "000000"})

    note_para = document.add_paragraph()
    note_para.paragraph_format.first_line_indent = Inches(0.2)
    note_para.paragraph_format.space_before = Pt(6)
    note_para.add_run(note)
    style_paragraph(note_para, font_size=7.5)
    try:
        document.save(path)
    except PermissionError:
        fallback = path.with_name(f"{path.stem}_updated{path.suffix}")
        document.save(fallback)
        print(f"Could not overwrite {path}; wrote {fallback} instead.")


def add_wide_metric_table(
    document: Document,
    data: pd.DataFrame,
    title_prefix: str,
    title: str,
    note: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(title_prefix)
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)

    display = data.drop(columns=["row_type"])
    table = document.add_table(rows=1, cols=len(display.columns))
    widths = []
    for column in display.columns:
        if column == "Validation set":
            widths.append(Inches(1.2))
        elif column == "Outcome":
            widths.append(Inches(0.85))
        elif column == "No. of protein markers":
            widths.append(Inches(0.75))
        else:
            widths.append(Inches(0.86))

    for i, column in enumerate(display.columns):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "BFBFBF")
        clear_borders(cell)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "12", "color": "000000"},
            bottom={"val": "single", "sz": "8", "color": "000000"},
        )
        set_cell_text(cell, column, 6.7, bold=True)

    for idx, row in data.iterrows():
        cells = table.add_row().cells
        values = display.loc[idx].tolist()
        is_section = row["row_type"] == "section"
        for i, value in enumerate(values):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            clear_borders(cells[i])
            if is_section:
                set_cell_shading(cells[i], "F2F2F2")
            set_cell_text(cells[i], value, 6.5, bold=is_section)

    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom={"val": "single", "sz": "12", "color": "000000"})

    note_para = document.add_paragraph()
    note_para.paragraph_format.first_line_indent = Inches(0.2)
    note_para.paragraph_format.space_before = Pt(6)
    note_para.add_run(note)
    style_paragraph(note_para, font_size=7.5)


def write_combined_supplementary_docx(
    cindex_external: pd.DataFrame, auc10: pd.DataFrame, path: Path
) -> None:
    document = Document()
    set_landscape(document)
    add_wide_metric_table(
        document,
        cindex_external,
        "Supplementary Table 3. ",
        "C-Index Across Non-European Validation Sets",
        "Values are C-index point estimates with 95% CIs in parentheses. Point estimates were recalculated using the full prediction files; 95% CIs were taken from bootstrap percentile intervals. Protein markers were selected separately for each outcome using LassoNet.",
    )
    document.add_page_break()
    add_wide_metric_table(
        document,
        auc10,
        "Supplementary Table 4. ",
        "10-Year Time-Dependent AUC Across Validation Sets",
        "Values are 10-year time-dependent AUC point estimates with 95% CIs in parentheses. 95% CIs were taken from bootstrap percentile intervals. Protein markers were selected separately for each outcome using LassoNet.",
    )
    try:
        document.save(path)
    except PermissionError:
        fallback = path.with_name(f"{path.stem}_updated{path.suffix}")
        document.save(fallback)
        print(f"Could not overwrite {path}; wrote {fallback} instead.")


def write_metric_docx(data: pd.DataFrame, path: Path, title_prefix: str, title: str, note: str) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(title_prefix)
    run.bold = True
    run.font.color.rgb = RGBColor(192, 0, 0) if title_prefix.startswith("Table") else RGBColor(0, 0, 0)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)

    display = data.drop(columns=["row_type"])
    table = document.add_table(rows=1, cols=len(display.columns))
    widths = [Inches(2.3), Inches(1.2), Inches(1.45), Inches(1.45), Inches(1.45)]
    for i, column in enumerate(display.columns):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        cell.text = column
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "BFBFBF")
        clear_borders(cell)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "12", "color": "000000"},
            bottom={"val": "single", "sz": "8", "color": "000000"},
        )
        for p in cell.paragraphs:
            style_paragraph(p, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for idx, row in data.iterrows():
        cells = table.add_row().cells
        values = display.loc[idx].tolist()
        is_section = row["row_type"] == "section"
        for i, value in enumerate(values):
            cells[i].width = widths[i]
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            clear_borders(cells[i])
            if is_section:
                set_cell_shading(cells[i], "F2F2F2")
            for p in cells[i].paragraphs:
                style_paragraph(p, bold=is_section)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom={"val": "single", "sz": "12", "color": "000000"})

    note_para = document.add_paragraph()
    note_para.paragraph_format.first_line_indent = Inches(0.2)
    note_para.paragraph_format.space_before = Pt(6)
    note_para.add_run(note)
    style_paragraph(note_para, font_size=8.2)
    document.save(path)


def write_readme() -> None:
    text = """3.2 Model performance evaluation 文件夹说明

本文件夹保存论文结果部分 3.2「Model performance evaluation」中模型区分度评价的可复现脚本和表格输出。

1. calculate_discrimination_metrics.py
   用于重新计算欧洲 hold-out set 中各模型的 C-index 和 10-year time-dependent AUC 点估计。
   点估计基于 saved_predictions 中保存的完整预测结果重新计算；95% CI 来自 summary.xlsx 中的 bootstrap 2.5%-97.5% 分位数。

2. Table2_cindex_european_holdout.xlsx / Table2_cindex_european_holdout.docx
   正文 Table 2。展示欧洲 hold-out set 中各模型的 C-index，保留 4 位小数。
   行按 Clinical predictors only、Protein markers only、Clinical predictors plus protein markers 分组；
   每个分组下按 Total CVD、ASCVD 和 HF 三个结局排列，列为模型。

3. Supplementary_Table_cindex_external_validation.xlsx / Supplementary_Table_cindex_external_validation.docx
   补充表 3。展示 Asian ancestry 和 Other ancestry 验证集中各模型的 C-index，保留 4 位小数。

4. Supplementary_Table_10yr_AUC_validation_sets.xlsx / Supplementary_Table_10yr_AUC_validation_sets.docx
   补充表。展示 European hold-out set、Asian ancestry 和 Other ancestry 验证集中各模型的
   10-year time-dependent AUC，保留 4 位小数。

主要计算口径：
- C-index 使用 risk_score 计算；risk_score 越高表示风险越高。
- 10-year time-dependent AUC 使用 prob_10yr 计算，时间点为 10 年。
- PREVENT equation 和 Refitted Cox model 使用规范化论文表述；其他模型名保留原始简称。
- No. of protein markers 来自各结局目录中的 lassonet_protein.csv。
"""
    (OUTPUT_DIR / "文件说明.txt").write_text(text, encoding="utf-8")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    cindex = build_metric_wide("C-index", validation="european")
    cindex_external = build_validation_wide("C-index", validations=["asian", "other"])
    auc10 = build_validation_wide("AUC10yr", validations=["european", "asian", "other"])

    cindex_out = OUTPUT_DIR / "Table2_cindex_european_holdout.xlsx"
    cindex_external_out = OUTPUT_DIR / "Supplementary_Table_cindex_external_validation.xlsx"
    auc_out = OUTPUT_DIR / "Supplementary_Table_10yr_AUC_validation_sets.xlsx"
    with pd.ExcelWriter(cindex_out, engine="openpyxl") as writer:
        cindex.drop(columns=["row_type"]).to_excel(writer, index=False, sheet_name="Table 2")
        ws = writer.sheets["Table 2"]
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = cell.alignment.copy(wrap_text=True, vertical="center")
    with pd.ExcelWriter(cindex_external_out, engine="openpyxl") as writer:
        cindex_external.drop(columns=["row_type"]).to_excel(
            writer, index=False, sheet_name="Supplementary C-index"
        )
        ws = writer.sheets["Supplementary C-index"]
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = cell.alignment.copy(wrap_text=True, vertical="center")
    with pd.ExcelWriter(auc_out, engine="openpyxl") as writer:
        auc10.drop(columns=["row_type"]).to_excel(writer, index=False, sheet_name="Supplementary AUC")
        ws = writer.sheets["Supplementary AUC"]
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = cell.alignment.copy(wrap_text=True, vertical="center")

    write_wide_metric_docx(
        cindex,
        OUTPUT_DIR / "Table2_cindex_european_holdout.docx",
        "Table 2. ",
        "Discrimination Performance in the European Hold-Out Set",
        "Values are C-index point estimates with 95% CIs in parentheses. Point estimates were recalculated using the full European hold-out prediction files; 95% CIs were taken from bootstrap percentile intervals. Protein markers were selected separately for each outcome using LassoNet.",
    )
    write_wide_metric_docx(
        cindex_external,
        OUTPUT_DIR / "Supplementary_Table_cindex_external_validation.docx",
        "Supplementary Table 3. ",
        "C-Index Across Non-European Validation Sets",
        "Values are C-index point estimates with 95% CIs in parentheses. Point estimates were recalculated using the full prediction files; 95% CIs were taken from bootstrap percentile intervals. Protein markers were selected separately for each outcome using LassoNet.",
    )
    write_combined_supplementary_docx(
        cindex_external,
        auc10,
        OUTPUT_DIR / "Supplementary_Table_10yr_AUC_validation_sets.docx",
    )
    write_readme()

    print(f"Wrote {cindex_out}")
    print(f"Wrote {cindex_external_out}")
    print(f"Wrote {auc_out}")
    print("C-index table:")
    print(cindex.drop(columns=["row_type"]).to_string(index=False))


if __name__ == "__main__":
    main()
