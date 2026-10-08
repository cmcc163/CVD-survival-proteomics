from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


OUT_DIR = Path(__file__).resolve().parent
SECTION_32 = Path(os.environ.get("CVD_PERFORMANCE_RESULTS", "results/performance"))
SHAP_DIR = Path(os.environ.get("CVD_SHAP_ROOT", "artifacts/shap/final"))

OUTCOMES = ["Total CVD", "ASCVD", "HF"]
OUTCOME_KEYS = {"Total CVD": "Total_CVD", "ASCVD": "ASCVD", "HF": "HF"}
PREDICTOR_SECTIONS = {
    "Clinical predictors only",
    "Protein markers only",
    "Clinical predictors plus protein markers",
}
BASE_ROWS = [
    ("PREVENT equation", "Clinical predictors only"),
    ("Refitted Cox model", "Clinical predictors only"),
    ("Refitted Cox model", "Protein markers only"),
    ("Refitted Cox model", "Clinical predictors plus protein markers"),
    ("SAINT", "Clinical predictors only"),
    ("SAINT", "Protein markers only"),
    ("SAINT", "Clinical predictors plus protein markers"),
]
REDUCED_MODEL = "SAINT reduced predictor panel"
REDUCED_SET = "Kneedle-selected clinical plus protein markers"


def flatten_discrimination_table(path: Path, sheet_name: str, validation: str | None = None) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name=sheet_name)
    rows = []
    current_set = None
    for _, row in raw.iterrows():
        first = row.iloc[0]
        outcome = row["Outcome"] if "Outcome" in raw.columns else None
        if isinstance(first, str) and first in PREDICTOR_SECTIONS:
            current_set = first
            continue
        if outcome not in OUTCOMES:
            continue
        if validation is not None and row.get("Validation set") != validation:
            continue
        for model, predictor_set in BASE_ROWS:
            if predictor_set == current_set:
                rows.append(
                    {
                        "Outcome": outcome,
                        "Model": model,
                        "Predictor set": predictor_set,
                        "Value": row[model],
                    }
                )
    return pd.DataFrame(rows)


def flatten_auc_table(path: Path) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name="Supplementary AUC")
    rows = []
    current_set = None
    for _, row in raw.iterrows():
        first = row.iloc[0]
        if isinstance(first, str) and first in PREDICTOR_SECTIONS:
            current_set = first
            continue
        if row.get("Validation set") != "European hold-out set" or row.get("Outcome") not in OUTCOMES:
            continue
        for model, predictor_set in BASE_ROWS:
            if predictor_set == current_set:
                rows.append(
                    {
                        "Outcome": row["Outcome"],
                        "Model": model,
                        "Predictor set": predictor_set,
                        "Value": row[model],
                    }
                )
    return pd.DataFrame(rows)


def previous_calibration(sheet_name: str) -> pd.DataFrame:
    path = SECTION_32 / "calibration_evaluation" / "Supplementary_Tables_calibration_metrics.xlsx"
    raw = pd.read_excel(path, sheet_name=sheet_name)
    keep = []
    for model, predictor_set in BASE_ROWS:
        keep.append(raw[
            raw["Outcome"].isin(OUTCOMES)
            & raw["Validation set"].eq("European hold-out set")
            & raw["Model"].eq(model)
            & raw["Predictor set"].eq(predictor_set)
        ])
    return pd.concat(keep, ignore_index=True)


def previous_net_benefit() -> pd.DataFrame:
    path = SECTION_32 / "clinical_utility_evaluation" / "Supplementary_Table_net_benefit_10yr.xlsx"
    raw = pd.read_excel(path, sheet_name="Supplementary table")
    keep = []
    for model, predictor_set in BASE_ROWS:
        keep.append(raw[
            raw["Outcome"].isin(OUTCOMES)
            & raw["Validation set"].eq("European hold-out set")
            & raw["Model"].eq(model)
            & raw["Predictor set"].eq(predictor_set)
        ])
    return pd.concat(keep, ignore_index=True)


def selected_features() -> pd.DataFrame:
    frames = []
    for outcome, key in OUTCOME_KEYS.items():
        path = SHAP_DIR / key / "Combined" / f"Selection_Results_{key}_Combined.xlsx"
        rank = pd.read_excel(path, sheet_name="1_Ensemble_SHAP_Rank")
        sel = rank[rank["Coarse_Kneedle"].astype(str).str.casefold().eq("yes")].copy()
        sel = sel[["Rank", "Feature", "Ensemble_Mean_SHAP_pct", "Cumulative_SHAP_pct", "Coarse_Kneedle"]]
        sel.insert(0, "Outcome", outcome)
        frames.append(sel)
    return pd.concat(frames, ignore_index=True)


def reduced_discrimination(metric: str) -> pd.DataFrame:
    df = pd.read_excel(OUT_DIR / "reduced_saint_discrimination_python.xlsx")
    col = "C-index formatted" if metric == "C-index" else "AUC 10yr formatted"
    return df.rename(columns={col: "Value"})[
        ["Outcome", "Model", "Predictor set", "Value"]
    ]


def reduced_calibration(sheet_name: str) -> pd.DataFrame:
    df = pd.read_excel(OUT_DIR / "reduced_saint_calibration_R.xlsx", sheet_name=sheet_name)
    return df


def reduced_net_benefit() -> pd.DataFrame:
    return pd.read_excel(OUT_DIR / "reduced_saint_net_benefit_R.xlsx", sheet_name="Supplementary table")


def metric_rows(metric_name: str, base: pd.DataFrame, reduced: pd.DataFrame) -> pd.DataFrame:
    combined = pd.concat([base, reduced], ignore_index=True)
    combined.insert(0, "Metric", metric_name)
    return combined


def build_integrated_summary() -> pd.DataFrame:
    cindex_base = flatten_discrimination_table(
        SECTION_32 / "discrimination_evaluation" / "Table2_cindex_european_holdout.xlsx",
        "Table 2",
    )
    auc_base = flatten_auc_table(
        SECTION_32 / "discrimination_evaluation" / "Supplementary_Table_10yr_AUC_validation_sets.xlsx"
    )
    cindex = metric_rows("C-index", cindex_base, reduced_discrimination("C-index"))
    auc = metric_rows("10-year AUC", auc_base, reduced_discrimination("AUC 10yr"))

    oe = previous_calibration("Supplementary Table 5")
    oe = pd.concat([oe, reduced_calibration("Supplementary Table 5")], ignore_index=True)
    oe.insert(0, "Metric", "O/E ratio")

    slope = previous_calibration("Supplementary Table 6")
    slope = pd.concat([slope, reduced_calibration("Supplementary Table 6")], ignore_index=True)
    slope.insert(0, "Metric", "Calibration slope")

    ici = previous_calibration("Supplementary Table 7")
    ici = pd.concat([ici, reduced_calibration("Supplementary Table 7")], ignore_index=True)
    ici.insert(0, "Metric", "ICI")

    nb = previous_net_benefit()
    nb = pd.concat([nb, reduced_net_benefit()], ignore_index=True)
    nb.insert(0, "Metric", "Net benefit")

    rows = []
    for outcome in OUTCOMES:
        for model, predictor_set in BASE_ROWS + [(REDUCED_MODEL, REDUCED_SET)]:
            row = {
                "Outcome": outcome,
                "Model": model,
                "Predictor set": predictor_set,
            }
            for frame, source_col, out_col in [
                (cindex, "Value", "C-index"),
                (auc, "Value", "10-year AUC"),
                (oe, "Overall", "O/E ratio overall"),
                (oe, "Low", "O/E ratio low"),
                (oe, "High", "O/E ratio high"),
                (slope, "Overall", "Calibration slope overall"),
                (slope, "Low", "Calibration slope low"),
                (slope, "High", "Calibration slope high"),
                (ici, "Overall", "ICI overall"),
                (ici, "Low", "ICI low"),
                (ici, "High", "ICI high"),
                (nb, "5%", "Net benefit 5%"),
                (nb, "7.5%", "Net benefit 7.5%"),
                (nb, "10%", "Net benefit 10%"),
            ]:
                hit = frame[
                    frame["Outcome"].eq(outcome)
                    & frame["Model"].eq(model)
                    & frame["Predictor set"].eq(predictor_set)
                ]
                row[out_col] = hit.iloc[0][source_col] if not hit.empty else "/"
            rows.append(row)
    return pd.DataFrame(rows)


def autosize(path: Path) -> None:
    wb = load_workbook(path)
    fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for col in ws.columns:
            width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 45)
    wb.save(path)


def main() -> None:
    summary = build_integrated_summary()
    selected = selected_features()
    out_xlsx = OUT_DIR / "European_holdout_reduced_panel_metrics.xlsx"
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        summary.to_excel(writer, index=False, sheet_name="Integrated summary")
        pd.read_excel(OUT_DIR / "reduced_saint_discrimination_python.xlsx").to_excel(
            writer, index=False, sheet_name="Reduced discrimination"
        )
        pd.read_excel(OUT_DIR / "reduced_saint_calibration_R.xlsx", sheet_name="Metrics_formatted").to_excel(
            writer, index=False, sheet_name="Reduced calibration"
        )
        pd.read_excel(OUT_DIR / "reduced_saint_net_benefit_R.xlsx", sheet_name="Supplementary table").to_excel(
            writer, index=False, sheet_name="Reduced net benefit"
        )
        selected.to_excel(writer, index=False, sheet_name="Kneedle selected predictors")
    autosize(out_xlsx)

    print(f"Saved {out_xlsx}")


if __name__ == "__main__":
    main()
