import os
from pathlib import Path

import pandas as pd
from sklearn.model_selection import train_test_split
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DATA_ROOT = Path(os.environ.get("CVD_DATA_ROOT", "data"))
INPUT_CSV = DATA_ROOT / "Total_CVD" / "followup_incident_20260122.csv"
OUTCOME_FILES = {
    "Total CVD": INPUT_CSV,
    "ASCVD": DATA_ROOT / "ASCVD" / "followup_incident_20260122.csv",
    "HF": DATA_ROOT / "HF" / "followup_incident_20260122.csv",
}
OUTPUT_DIR = Path(__file__).resolve().parent
EXCEL_OUT = OUTPUT_DIR / "table1_total_cvd_participant_characteristics.xlsx"
WORD_OUT = OUTPUT_DIR / "Table1_total_CVD_participant_characteristics.docx"
SUPP_EXCEL_OUT = OUTPUT_DIR / "Supplementary_Table_outcome_followup_and_ancestry_summary.xlsx"
SUPP_WORD_OUT = OUTPUT_DIR / "Supplementary_Table_outcome_followup_and_ancestry_summary.docx"
SUPP_BASELINE_EXCEL_OUT = OUTPUT_DIR / "Supplementary_Table_dataset_baseline_characteristics_by_outcome.xlsx"
SUPP_BASELINE_WORD_OUT = OUTPUT_DIR / "Supplementary_Table_dataset_baseline_characteristics_by_outcome.docx"

REQUIRED_COLUMNS = [
    "eid",
    "Is_Incident",
    "Ethnic",
    "age",
    "sex",
    "ever_smoked",
    "Diabetes_baseline_1",
    "Cholesterol_treatment",
    "hdl_cholesterol",
    "total_cholesterol",
    "hypertension_treatment",
    "average_SBP",
    "eGFR_SCysC",
    "BMI",
]

FOLLOWUP_COLUMNS = ["eid", "Ethnic", "baseline_date", "Event_Date", "Is_Incident"]

EUROPEAN_CODES = {1, 1001, 1002, 1003}
ASIAN_CODES = {3, 3001, 3002, 3003, 3004, 5}


def format_int(n: int) -> str:
    return f"{int(n):,}".replace(",", " ")


def format_mean_sd(series: pd.Series) -> str:
    values = pd.to_numeric(series, errors="coerce").dropna()
    if values.empty:
        return ""
    return f"{values.mean():.2f}\u00b1{values.std(ddof=1):.2f}"


def format_n_pct(mask: pd.Series, denominator: int) -> str:
    n = int(mask.sum())
    pct = 100 * n / denominator if denominator else 0
    return f"{format_int(n)} ({pct:.1f})"


def ethnicity_group(value) -> str:
    if pd.isna(value):
        return "Missing"
    try:
        code = int(value)
    except (TypeError, ValueError):
        return "Missing"
    if code in EUROPEAN_CODES:
        return "European"
    if code in ASIAN_CODES:
        return "Asian"
    return "Other"


def summarize_group(df: pd.DataFrame, label: str, subset: pd.DataFrame) -> dict:
    denominator = len(subset)
    return {
        "group": label,
        "n": denominator,
        "Age, y": format_mean_sd(subset["age"]),
        "Male sex": format_n_pct(subset["sex"].eq(1), denominator),
        "Ever smoked": format_n_pct(subset["ever_smoked"].eq(1), denominator),
        "Diabetes": format_n_pct(subset["Diabetes_baseline_1"].eq(1), denominator),
        "Lipid-lowering medication": format_n_pct(
            subset["Cholesterol_treatment"].eq(1), denominator
        ),
        "HDL cholesterol, mmol/L": format_mean_sd(subset["hdl_cholesterol"]),
        "Non-HDL cholesterol, mmol/L": format_mean_sd(subset["non_hdl_cholesterol"]),
        "Antihypertensive medication": format_n_pct(
            subset["hypertension_treatment"].eq(1), denominator
        ),
        "Systolic blood pressure, mm Hg": format_mean_sd(subset["average_SBP"]),
        "eGFR, mL/min/1.73 m\u00b2": format_mean_sd(subset["eGFR_SCysC"]),
        "BMI, kg/m\u00b2": format_mean_sd(subset["BMI"]),
    }


def make_summary(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["non_hdl_cholesterol"] = df["total_cholesterol"] - df["hdl_cholesterol"]
    df["ethnicity_group"] = df["Ethnic"].map(ethnicity_group)

    groups = [
        ("Overall", df),
        ("No incident CVD", df[df["Is_Incident"].eq(0)]),
        ("Incident CVD", df[df["Is_Incident"].eq(1)]),
    ]
    headers = [f"{label} (n={format_int(len(subset))})" for label, subset in groups]

    rows = []

    def add_row(characteristic: str, key: str):
        rows.append(
            {
                "Characteristic": characteristic,
                headers[0]: summary_by_group[0][key],
                headers[1]: summary_by_group[1][key],
                headers[2]: summary_by_group[2][key],
            }
        )

    summary_by_group = [summarize_group(df, label, subset) for label, subset in groups]

    add_row("Age, y", "Age, y")
    add_row("Male sex", "Male sex")
    rows.append(
        {
            "Characteristic": "Genetic ancestry group",
            headers[0]: "",
            headers[1]: "",
            headers[2]: "",
        }
    )
    for category in ["European", "Asian", "Other"]:
        row = {"Characteristic": f"   {category}"}
        for header, (_, subset) in zip(headers, groups):
            row[header] = format_n_pct(subset["ethnicity_group"].eq(category), len(subset))
        rows.append(row)
    add_row("Ever smoked", "Ever smoked")
    add_row("Diabetes", "Diabetes")
    add_row("Lipid-lowering medication", "Lipid-lowering medication")
    add_row("HDL cholesterol, mmol/L", "HDL cholesterol, mmol/L")
    add_row("Non-HDL cholesterol, mmol/L", "Non-HDL cholesterol, mmol/L")
    add_row("Antihypertensive medication", "Antihypertensive medication")
    add_row("Systolic blood pressure, mm Hg", "Systolic blood pressure, mm Hg")
    add_row("eGFR, mL/min/1.73 m\u00b2", "eGFR, mL/min/1.73 m\u00b2")
    add_row("BMI, kg/m\u00b2", "BMI, kg/m\u00b2")

    return pd.DataFrame(rows), headers


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn(f"w:{key}"), str(value))


def clear_borders(cell):
    nil = {"val": "nil"}
    set_cell_border(cell, top=nil, bottom=nil, left=nil, right=nil, insideH=nil, insideV=nil)


def add_three_line_table(document: Document, data: pd.DataFrame, widths: list[Inches]) -> None:
    table = document.add_table(rows=1, cols=len(data.columns))
    table.autofit = False
    header_cells = table.rows[0].cells
    for i, text in enumerate(data.columns):
        cell = header_cells[i]
        cell.width = widths[i]
        cell.text = str(text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "BFBFBF")
        clear_borders(cell)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "12", "color": "000000"},
            bottom={"val": "single", "sz": "8", "color": "000000"},
        )
        for p in cell.paragraphs:
            style_paragraph(p, font_size=8.2, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for _, row in data.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row.tolist()):
            cells[i].width = widths[i]
            cells[i].text = "" if pd.isna(value) else str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            clear_borders(cells[i])
            for p in cells[i].paragraphs:
                style_paragraph(p, font_size=8.2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for cell in table.rows[-1].cells:
        set_cell_border(cell, bottom={"val": "single", "sz": "12", "color": "000000"})
    return table


def style_paragraph(paragraph, font_size=9, bold=False, color=None):
    for run in paragraph.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(font_size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)


def write_word(summary: pd.DataFrame, headers: list[str]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Table 1. ")
    run.bold = True
    run.font.color.rgb = RGBColor(192, 0, 0)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)
    run = title.add_run(
        "Participant Characteristics at Baseline by Incident Total Cardiovascular Disease During Follow-Up"
    )
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)

    widths = [Inches(2.55), Inches(1.85), Inches(1.85), Inches(1.85)]
    table_data = summary.rename(columns={"Characteristic": "Characteristic"})[
        ["Characteristic", *headers]
    ].fillna("")
    table = add_three_line_table(document, table_data, widths)
    for row in table.rows:
        is_section = row.cells[0].text == "Genetic ancestry group"
        if is_section:
            for cell in row.cells:
                set_cell_shading(cell, "F2F2F2")
                for p in cell.paragraphs:
                    style_paragraph(p, font_size=8.2, bold=True)

    note = document.add_paragraph()
    note.paragraph_format.first_line_indent = Inches(0.2)
    note.paragraph_format.space_before = Pt(6)
    note.add_run(
        "Values are mean\u00b1SD or n (%) unless otherwise specified. BMI indicates body mass index; "
        "CVD, cardiovascular disease; eGFR, estimated glomerular filtration rate based on serum creatinine and cystatin C; "
        "HDL, high-density lipoprotein; and SBP, systolic blood pressure. "
        "Non-HDL cholesterol was calculated as total cholesterol minus HDL cholesterol."
    )
    for p in [note]:
        style_paragraph(p, font_size=8.5)

    document.save(WORD_OUT)


def load_followup(path: Path) -> pd.DataFrame:
    data = pd.read_csv(path, usecols=FOLLOWUP_COLUMNS)
    data["baseline_date"] = pd.to_datetime(data["baseline_date"], errors="coerce")
    data["Event_Date"] = pd.to_datetime(data["Event_Date"], errors="coerce")
    data["time_days"] = (data["Event_Date"] - data["baseline_date"]).dt.days
    data = data.dropna(subset=["time_days", "Is_Incident", "Ethnic"])
    data = data[data["time_days"] > 0].copy()
    data["time_years"] = data["time_days"] / 365.25
    data["ancestry_group"] = data["Ethnic"].map(ethnicity_group)
    return data


def bootstrap_median_ci(values: pd.Series, n_boot: int = 2000, seed: int = 221) -> tuple[float, float, float]:
    values = pd.to_numeric(values, errors="coerce").dropna().to_numpy()
    median = float(pd.Series(values).median())
    if len(values) < 2:
        return median, median, median
    samples = []
    rng = pd.Series(range(n_boot)).sample(n=n_boot, random_state=seed).to_numpy()
    import numpy as np

    generator = np.random.default_rng(seed)
    for _ in rng:
        samples.append(float(np.median(generator.choice(values, size=len(values), replace=True))))
    lower, upper = np.percentile(samples, [2.5, 97.5])
    return median, float(lower), float(upper)


def poisson_rate_ci(events: int, person_years: float, multiplier: int = 1000) -> tuple[float, float, float]:
    import math

    rate = events / person_years * multiplier if person_years else 0
    if events == 0 or person_years == 0:
        return rate, 0.0, 0.0
    se_log = 1 / math.sqrt(events)
    lower = math.exp(math.log(rate) - 1.96 * se_log)
    upper = math.exp(math.log(rate) + 1.96 * se_log)
    return rate, lower, upper


def format_est_ci(estimate: float, lower: float, upper: float, digits: int = 2) -> str:
    return f"{estimate:.{digits}f} ({lower:.{digits}f}-{upper:.{digits}f})"


def reverse_km_median_ci(time_years: pd.Series, incident_event: pd.Series) -> tuple[float, float, float]:
    from lifelines import KaplanMeierFitter
    import numpy as np

    data = pd.DataFrame(
        {
            "time": pd.to_numeric(time_years, errors="coerce"),
            "censoring_event": 1 - pd.to_numeric(incident_event, errors="coerce"),
        }
    ).dropna()
    kmf = KaplanMeierFitter(alpha=0.05)
    kmf.fit(data["time"], event_observed=data["censoring_event"])

    median = float(kmf.median_survival_time_)
    ci = kmf.confidence_interval_survival_function_
    survival = kmf.survival_function_[kmf.survival_function_.columns[0]]
    times = survival.index.to_numpy(dtype=float)
    lower_survival = ci.iloc[:, 0].to_numpy(dtype=float)
    upper_survival = ci.iloc[:, 1].to_numpy(dtype=float)

    def first_crossing(bound: np.ndarray) -> float:
        crossed = np.where(bound <= 0.5)[0]
        if len(crossed) == 0:
            return float("nan")
        return float(times[crossed[0]])

    lower_median = first_crossing(lower_survival)
    upper_median = first_crossing(upper_survival)
    if np.isnan(lower_median):
        lower_median = median
    if np.isnan(upper_median):
        upper_median = median
    lower_median, upper_median = sorted([lower_median, upper_median])
    return median, lower_median, upper_median


def make_outcome_followup_summary() -> pd.DataFrame:
    rows = []
    for outcome, path in OUTCOME_FILES.items():
        data = load_followup(path)
        eur = data[data["ancestry_group"].eq("European")]
        if eur["Is_Incident"].nunique() > 1:
            eur_dev_idx, eur_holdout_idx = train_test_split(
                eur.index,
                test_size=0.1,
                random_state=221,
                stratify=eur["Is_Incident"],
            )
        else:
            eur_dev_idx, eur_holdout_idx = train_test_split(
                eur.index,
                test_size=0.1,
                random_state=221,
            )
        median, lower, upper = reverse_km_median_ci(data["time_years"], data["Is_Incident"])
        events = int(data["Is_Incident"].sum())
        rows.append(
            {
                "Outcome": outcome,
                "Participants, n": format_int(len(data)),
                "Events, n": format_int(events),
                "Median follow-up, y (95% CI)": format_est_ci(median, lower, upper),
                "European development set, n": format_int(len(eur_dev_idx)),
                "European hold-out set, n": format_int(len(eur_holdout_idx)),
                "Asian ancestry, n": format_int(data["ancestry_group"].eq("Asian").sum()),
                "Other ancestry, n": format_int(data["ancestry_group"].eq("Other").sum()),
            }
        )
    return pd.DataFrame(rows)


def load_baseline_for_outcome(path: Path) -> pd.DataFrame:
    columns = list(dict.fromkeys(REQUIRED_COLUMNS + FOLLOWUP_COLUMNS))
    data = pd.read_csv(path, usecols=columns)
    data["total_cholesterol"] = pd.to_numeric(data["total_cholesterol"], errors="coerce")
    data["hdl_cholesterol"] = pd.to_numeric(data["hdl_cholesterol"], errors="coerce")
    data["non_hdl_cholesterol"] = data["total_cholesterol"] - data["hdl_cholesterol"]
    data["baseline_date"] = pd.to_datetime(data["baseline_date"], errors="coerce")
    data["Event_Date"] = pd.to_datetime(data["Event_Date"], errors="coerce")
    data["time_days"] = (data["Event_Date"] - data["baseline_date"]).dt.days
    data = data.dropna(subset=["time_days", "Is_Incident", "Ethnic"])
    data = data[data["time_days"] > 0].copy()
    data["ancestry_group"] = data["Ethnic"].map(ethnicity_group)
    return data


def split_outcome_datasets(data: pd.DataFrame) -> dict[str, pd.DataFrame]:
    eur = data[data["ancestry_group"].eq("European")]
    asian = data[data["ancestry_group"].eq("Asian")]
    other = data[data["ancestry_group"].eq("Other")]
    if eur["Is_Incident"].nunique() > 1:
        eur_dev_idx, eur_holdout_idx = train_test_split(
            eur.index,
            test_size=0.1,
            random_state=221,
            stratify=eur["Is_Incident"],
        )
    else:
        eur_dev_idx, eur_holdout_idx = train_test_split(
            eur.index,
            test_size=0.1,
            random_state=221,
        )
    return {
        "European development set": data.loc[eur_dev_idx],
        "European hold-out set": data.loc[eur_holdout_idx],
        "Asian ancestry": asian,
        "Other ancestry": other,
    }


def summarize_baseline_dataset(subset: pd.DataFrame) -> dict[str, str]:
    denominator = len(subset)
    return {
        "Participants, n": format_int(denominator),
        "Events, n": format_int(int(subset["Is_Incident"].sum())),
        "Age, y": format_mean_sd(subset["age"]),
        "Male sex": format_n_pct(subset["sex"].eq(1), denominator),
        "Ever smoked": format_n_pct(subset["ever_smoked"].eq(1), denominator),
        "Diabetes": format_n_pct(subset["Diabetes_baseline_1"].eq(1), denominator),
        "Lipid-lowering medication": format_n_pct(
            subset["Cholesterol_treatment"].eq(1), denominator
        ),
        "HDL cholesterol, mmol/L": format_mean_sd(subset["hdl_cholesterol"]),
        "Non-HDL cholesterol, mmol/L": format_mean_sd(subset["non_hdl_cholesterol"]),
        "Antihypertensive medication": format_n_pct(
            subset["hypertension_treatment"].eq(1), denominator
        ),
        "Systolic blood pressure, mm Hg": format_mean_sd(subset["average_SBP"]),
        "eGFR, mL/min/1.73 m\u00b2": format_mean_sd(subset["eGFR_SCysC"]),
        "BMI, kg/m\u00b2": format_mean_sd(subset["BMI"]),
    }


def make_dataset_baseline_summary() -> pd.DataFrame:
    outcome_summaries = {}
    dataset_order = [
        "European development set",
        "European hold-out set",
        "Asian ancestry",
        "Other ancestry",
    ]
    characteristic_order = [
        "Participants, n",
        "Events, n",
        "Age, y",
        "Male sex",
        "Ever smoked",
        "Diabetes",
        "Lipid-lowering medication",
        "HDL cholesterol, mmol/L",
        "Non-HDL cholesterol, mmol/L",
        "Antihypertensive medication",
        "Systolic blood pressure, mm Hg",
        "eGFR, mL/min/1.73 m\u00b2",
        "BMI, kg/m\u00b2",
    ]
    for outcome, path in OUTCOME_FILES.items():
        data = load_baseline_for_outcome(path)
        splits = split_outcome_datasets(data)
        outcome_summaries[outcome] = {
            dataset: summarize_baseline_dataset(subset) for dataset, subset in splits.items()
        }

    rows = []
    for outcome in OUTCOME_FILES.keys():
        for characteristic in characteristic_order:
            row = {"Outcome": outcome, "Characteristic": characteristic}
            for dataset in dataset_order:
                row[dataset] = outcome_summaries[outcome][dataset][characteristic]
            rows.append(row)
    return pd.DataFrame(rows)


def write_supplementary_word(summary: pd.DataFrame) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Supplementary Table 1. ")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)
    run = title.add_run(
        "Outcome, Follow-Up, Incidence, and Ancestry-Specific Sample Sizes"
    )
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)

    widths = [
        Inches(0.75),
        Inches(0.75),
        Inches(0.65),
        Inches(1.1),
        Inches(0.95),
        Inches(0.9),
        Inches(0.75),
        Inches(0.75),
    ]
    add_three_line_table(document, summary, widths)

    note = document.add_paragraph()
    note.paragraph_format.first_line_indent = Inches(0.2)
    note.paragraph_format.space_before = Pt(6)
    note.add_run(
        "Follow-up time was calculated as Event_Date minus baseline_date after excluding participants with missing or nonpositive follow-up. "
        "Median follow-up and 95% confidence intervals were estimated using the reverse Kaplan-Meier method. "
        "European development and hold-out sets reproduce the modeling split among European participants using test_size=0.1, random_state=221, and event-stratified sampling."
    )
    style_paragraph(note, font_size=8.2)

    document.save(SUPP_WORD_OUT)


def write_dataset_baseline_word(summary: pd.DataFrame) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Supplementary Table 2. ")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)
    run = title.add_run(
        "Baseline Characteristics by Modeling Dataset and Cardiovascular Outcome"
    )
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)

    display = summary.copy()
    widths = [Inches(0.85), Inches(1.9), Inches(1.45), Inches(1.45), Inches(1.25), Inches(1.25)]
    table = add_three_line_table(document, display, widths)
    previous_outcome = None
    for row in table.rows[1:]:
        outcome = row.cells[0].text
        if outcome == previous_outcome:
            row.cells[0].text = ""
        else:
            previous_outcome = outcome
            for cell in row.cells:
                set_cell_border(cell, top={"val": "single", "sz": "4", "color": "808080"})

    note = document.add_paragraph()
    note.paragraph_format.first_line_indent = Inches(0.2)
    note.paragraph_format.space_before = Pt(6)
    note.add_run(
        "Values are mean\u00b1SD or n (%) unless otherwise specified. Dataset definitions reproduce the modeling split for each outcome. "
        "BMI indicates body mass index; eGFR, estimated glomerular filtration rate based on serum creatinine and cystatin C; HDL, high-density lipoprotein; and SBP, systolic blood pressure."
    )
    style_paragraph(note, font_size=8.2)

    document.save(SUPP_BASELINE_WORD_OUT)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    df = pd.read_csv(INPUT_CSV, usecols=REQUIRED_COLUMNS)
    summary, headers = make_summary(df)

    supp_summary = make_outcome_followup_summary()
    supp_baseline_summary = make_dataset_baseline_summary()

    with pd.ExcelWriter(EXCEL_OUT, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Table 1", index=False)
        df.assign(
            non_hdl_cholesterol=df["total_cholesterol"] - df["hdl_cholesterol"],
            ethnicity_group=df["Ethnic"].map(ethnicity_group),
        )[REQUIRED_COLUMNS + ["non_hdl_cholesterol", "ethnicity_group"]].to_excel(
            writer, sheet_name="Analysis data", index=False
        )

    with pd.ExcelWriter(SUPP_EXCEL_OUT, engine="openpyxl") as writer:
        supp_summary.to_excel(writer, sheet_name="Supplementary Table 1", index=False)

    with pd.ExcelWriter(SUPP_BASELINE_EXCEL_OUT, engine="openpyxl") as writer:
        supp_baseline_summary.to_excel(writer, sheet_name="Supplementary Table 2", index=False)

    write_word(summary, headers)
    write_supplementary_word(supp_summary)
    write_dataset_baseline_word(supp_baseline_summary)
    print(f"Wrote {EXCEL_OUT}")
    print(f"Wrote {WORD_OUT}")
    print(f"Wrote {SUPP_EXCEL_OUT}")
    print(f"Wrote {SUPP_WORD_OUT}")
    print(f"Wrote {SUPP_BASELINE_EXCEL_OUT}")
    print(f"Wrote {SUPP_BASELINE_WORD_OUT}")


if __name__ == "__main__":
    main()
