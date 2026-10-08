import os
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DATA_ROOT = Path(os.environ.get("CVD_DATA_ROOT", "data"))
OUTCOME_FILES = {
    label: DATA_ROOT / key / "followup_incident_20260122.csv"
    for label, key in (("Total CVD", "Total_CVD"), ("ASCVD", "ASCVD"), ("HF", "HF"))
}

OUTPUT_FILE = Path(__file__).resolve().parent / "Outcome_followup_person_years.xlsx"
USE_COLUMNS = ["eid", "Ethnic", "baseline_date", "Event_Date", "Is_Incident"]


def calculate_person_years(outcome: str, path: Path) -> dict:
    data = pd.read_csv(path, usecols=USE_COLUMNS)
    data["baseline_date"] = pd.to_datetime(data["baseline_date"], errors="coerce")
    data["Event_Date"] = pd.to_datetime(data["Event_Date"], errors="coerce")
    data["followup_days"] = (data["Event_Date"] - data["baseline_date"]).dt.days

    valid = (
        data["followup_days"].notna()
        & data["Is_Incident"].notna()
        & data["Ethnic"].notna()
        & data["followup_days"].gt(0)
    )
    analysis = data.loc[valid].copy()
    analysis["followup_years"] = analysis["followup_days"] / 365.25

    return {
        "Outcome": outcome,
        "Source file": str(path),
        "Raw records, n": len(data),
        "Included participants, n": len(analysis),
        "Excluded records, n": int((~valid).sum()),
        "Incident events, n": int(
            pd.to_numeric(analysis["Is_Incident"], errors="coerce").sum()
        ),
        "Total follow-up, person-years": analysis["followup_years"].sum(),
        "Mean follow-up, years": analysis["followup_years"].mean(),
        "Median individual follow-up, years": analysis["followup_years"].median(),
        "Minimum follow-up, days": int(analysis["followup_days"].min()),
        "Maximum follow-up, days": int(analysis["followup_days"].max()),
        "Unique participants, n": analysis["eid"].nunique(),
        "Duplicate eid records, n": int(analysis["eid"].duplicated().sum()),
    }


def format_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in worksheet.columns:
            letter = get_column_letter(column[0].column)
            max_length = max(len(str(cell.value or "")) for cell in column)
            worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 65)

    summary = workbook["Person-years summary"]
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 7).number_format = '#,##0.00'
        summary.cell(row, 8).number_format = '0.00'
        summary.cell(row, 9).number_format = '0.00'

    workbook.save(path)


def main() -> None:
    summary = pd.DataFrame(
        calculate_person_years(outcome, path)
        for outcome, path in OUTCOME_FILES.items()
    )
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Person-years summary", index=False)
    format_workbook(OUTPUT_FILE)
    print(summary[["Outcome", "Included participants, n", "Incident events, n", "Total follow-up, person-years"]].to_string(index=False))
    print(f"Wrote {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
